"""Generates the synthetic demo attachments used by `MailStudio.exe --demo` (no real data).

Usage:  python make_demo_files.py <output-folder>
Needs:  ezdxf, matplotlib, openpyxl, python-docx, python-pptx  (STEP is created by make_demo_step.py with OCP)
"""
import math, os, sys
import ezdxf
from ezdxf.enums import TextEntityAlignment

out = sys.argv[1] if len(sys.argv) > 1 else 'demo'
os.makedirs(out, exist_ok=True)


# ----------------------------------------------------------------------------- DXF: flange drawing
def flange_dxf(path):
    doc = ezdxf.new('R2018', setup=True)
    doc.units = ezdxf.units.MM
    for name, color, lt in [('OUTLINE', 7, 'CONTINUOUS'), ('CENTER', 1, 'CENTER'), ('HIDDEN', 8, 'DASHED'),
                            ('DIMS', 3, 'CONTINUOUS'), ('HATCH', 252, 'CONTINUOUS'), ('TEXT', 5, 'CONTINUOUS'), ('TITLE', 7, 'CONTINUOUS')]:
        doc.layers.add(name, color=color, linetype=lt)
    msp = doc.modelspace()
    cx, cy = 150, 195
    msp.add_circle((cx, cy), 110, dxfattribs={'layer': 'OUTLINE', 'lineweight': 50})
    msp.add_circle((cx, cy), 55, dxfattribs={'layer': 'OUTLINE', 'lineweight': 50})
    msp.add_circle((cx, cy), 90, dxfattribs={'layer': 'CENTER'})
    for i in range(8):
        a = math.radians(22.5 + i * 45)
        hx, hy = cx + 90 * math.cos(a), cy + 90 * math.sin(a)
        msp.add_circle((hx, hy), 9, dxfattribs={'layer': 'OUTLINE', 'lineweight': 35})
        msp.add_line((hx - 13, hy), (hx + 13, hy), dxfattribs={'layer': 'CENTER'})
        msp.add_line((hx, hy - 13), (hx, hy + 13), dxfattribs={'layer': 'CENTER'})
    msp.add_line((cx - 125, cy), (cx + 125, cy), dxfattribs={'layer': 'CENTER'})
    msp.add_line((cx, cy - 125), (cx, cy + 125), dxfattribs={'layer': 'CENTER'})
    sx = 330
    outline = [(sx, cy - 110), (sx + 24, cy - 110), (sx + 24, cy - 70), (sx + 60, cy - 66), (sx + 60, cy - 55), (sx, cy - 55)]
    for sign in (1, -1):
        pts = [(x, cy + sign * (y - cy)) for x, y in outline]
        msp.add_lwpolyline(pts, close=True, dxfattribs={'layer': 'OUTLINE', 'lineweight': 50})
        h = msp.add_hatch(color=252, dxfattribs={'layer': 'HATCH'})
        h.set_pattern_fill('ANSI31', scale=1.2)
        h.paths.add_polyline_path(pts, is_closed=True)
    msp.add_line((sx - 10, cy), (sx + 70, cy), dxfattribs={'layer': 'CENTER'})
    for y in (cy + 90, cy - 90):
        msp.add_line((sx - 4, y), (sx + 28, y), dxfattribs={'layer': 'CENTER'})
    o = {'dimtxt': 5, 'dimasz': 4, 'dimlfac': 1}
    msp.add_diameter_dim(center=(cx, cy), radius=110, angle=135, dimstyle='EZ_MM_RADIUS', override=o, dxfattribs={'layer': 'DIMS'}).render()
    msp.add_diameter_dim(center=(cx, cy), radius=55, angle=30, dimstyle='EZ_MM_RADIUS', override=o, dxfattribs={'layer': 'DIMS'}).render()
    msp.add_linear_dim(base=(sx + 75, cy), p1=(sx + 24, cy + 110), p2=(sx + 24, cy - 110), angle=90, dimstyle='EZDXF', override=dict(o, dimexe=2), dxfattribs={'layer': 'DIMS'}).render()
    msp.add_linear_dim(base=(sx, cy - 130), p1=(sx, cy - 110), p2=(sx + 60, cy - 66), dimstyle='EZDXF', override=o, dxfattribs={'layer': 'DIMS'}).render()
    msp.add_text('8x Ø18 THRU EQ. SP. ON PCD 180', height=5, dxfattribs={'layer': 'TEXT'}).set_placement((cx - 80, cy - 140), align=TextEntityAlignment.LEFT)
    msp.add_text('SECTION A-A', height=6, dxfattribs={'layer': 'TEXT'}).set_placement((sx - 5, cy + 125), align=TextEntityAlignment.LEFT)
    msp.add_text('Ra 3.2', height=4, dxfattribs={'layer': 'TEXT'}).set_placement((sx + 30, cy + 80), align=TextEntityAlignment.LEFT)
    msp.add_lwpolyline([(0, 0), (480, 0), (480, 345), (0, 345)], close=True, dxfattribs={'layer': 'TITLE', 'lineweight': 70})
    msp.add_lwpolyline([(300, 0), (480, 0), (480, 50), (300, 50)], close=True, dxfattribs={'layer': 'TITLE', 'lineweight': 50})
    for y in (16, 33):
        msp.add_line((300, y), (480, y), dxfattribs={'layer': 'TITLE'})
    msp.add_line((390, 0), (390, 33), dxfattribs={'layer': 'TITLE'})
    for txt, x, y, h in [('FLANGE DN100 PN16', 305, 38, 7), ('DWG NO. NW-1042', 305, 21, 5), ('MAT. S235JR', 395, 21, 5),
                         ('SCALE 1:2   SHEET 1/1', 305, 5, 4.5), ('NORTHWIND MACHINING (DEMO)', 395, 5, 3.6)]:
        msp.add_text(txt, height=h, dxfattribs={'layer': 'TITLE'}).set_placement((x, y), align=TextEntityAlignment.LEFT)
    doc.saveas(path)
    return doc


def dxf_to_pdf(doc, path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from ezdxf.addons.drawing import RenderContext, Frontend
    from ezdxf.addons.drawing.matplotlib import MatplotlibBackend
    from ezdxf.addons.drawing.config import Configuration, BackgroundPolicy, ColorPolicy
    fig = plt.figure(figsize=(16.54, 11.69))
    ax = fig.add_axes([0.02, 0.02, 0.96, 0.96])
    cfg = Configuration(background_policy=BackgroundPolicy.WHITE, color_policy=ColorPolicy.BLACK)
    Frontend(RenderContext(doc), MatplotlibBackend(ax), config=cfg).draw_layout(doc.modelspace(), finalize=True)
    fig.savefig(path)
    plt.close(fig)


def schedule_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = 'Week 43'
    ws.append(['Part no.', 'Description', 'Material', 'Qty', 'Unit price (EUR)', 'Total (EUR)', 'Due date', 'Status'])
    rows = [('NW-1042', 'Flange DN100 PN16', 'S235JR', 120, 38.5, '2026-10-23', 'In production'),
            ('NW-1043', 'Flange DN80 PN16', 'S235JR', 60, 31.2, '2026-10-23', 'In production'),
            ('CA-2201', 'Bracket rev C', 'AlMg3', 40, 22.9, '2026-10-27', 'CAM ready'),
            ('CA-2202', 'Spacer Ø40x12', '1.4301', 250, 3.4, '2026-10-29', 'Waiting material'),
            ('CA-2210', 'Shaft Ø25 h7', '42CrMo4', 30, 64.0, '2026-11-02', 'Quoted'),
            ('NW-1050', 'Cover plate 300x200', 'S355J2', 18, 47.8, '2026-11-05', 'Quoted')]
    for i, r in enumerate(rows, start=2):
        ws.append([r[0], r[1], r[2], r[3], r[4], f'=D{i}*E{i}', r[5], r[6]])
    n = len(rows) + 1
    ws.append(['', 'Total', '', f'=SUM(D2:D{n})', '', f'=SUM(F2:F{n})', '', ''])
    thin = Side(style='thin', color='BFBFBF')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79'); c.alignment = Alignment(horizontal='center')
    for row in ws.iter_rows(min_row=2, max_row=n + 1):
        for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c in ws[n + 1]: c.font = Font(bold=True)
    for col, w in zip('ABCDEFGH', (12, 26, 12, 8, 16, 14, 13, 18)): ws.column_dimensions[col].width = w
    for r in range(2, n + 2):
        ws[f'E{r}'].number_format = '#,##0.00'; ws[f'F{r}'].number_format = '#,##0.00'
    wb.save(path)


def report_docx(path):
    from docx import Document
    from docx.shared import RGBColor
    d = Document()
    d.add_heading('First Article Inspection Report', 0)
    p = d.add_paragraph(); p.add_run('Part: ').bold = True; p.add_run('NW-1042 Flange DN100 PN16    ')
    p.add_run('Drawing: ').bold = True; p.add_run('NW-1042 rev B')
    d.add_heading('1. Measured characteristics', level=1)
    t = d.add_table(rows=1, cols=5); t.style = 'Light Grid Accent 1'
    for c, h in zip(t.rows[0].cells, ('#', 'Characteristic', 'Nominal', 'Tolerance', 'Measured')): c.text = h
    for row in [('1', 'Outer diameter', 'Ø220', '±0.2', '220.06'), ('2', 'Bore', 'Ø110', 'H7', '110.012'),
                ('3', 'Pitch circle', 'Ø180', '±0.1', '180.03'), ('4', 'Thickness', '24', '±0.1', '23.97'), ('5', 'Hole size (8x)', 'Ø18', '+0.2/0', '18.08')]:
        for c, v in zip(t.add_row().cells, row): c.text = v
    d.add_heading('2. Result', level=1)
    r = d.add_paragraph().add_run('All characteristics within tolerance. Released for serial production.')
    r.font.color.rgb = RGBColor(0x1E, 0x7B, 0x34); r.bold = True
    d.add_paragraph('Inspector: J. Demo        Date: 2026-10-16')
    d.save(path)


def kickoff_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = 'Project Kickoff — Bracket rev C'
    s.placeholders[1].text = 'Contoso Aero · Northwind Machining (demo data)'
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = 'Scope'
    s.placeholders[1].text = 'Material: AlMg3, 40 pcs\nTolerances: ISO 2768-mK\nSurface: anodized, clear\nFirst article inspection required'
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = 'Timeline'
    s.placeholders[1].text = 'Week 43: CAM and fixture\nWeek 44: first article\nWeek 45: serial production\nWeek 46: delivery'
    prs.save(path)


def op20_nc(path):
    open(path, 'w', encoding='ascii', newline='\r\n').write('%\nO2020 (BRACKET REV C - OP20)\n(DEMO PROGRAM - NOT FOR PRODUCTION)\nG21 G17 G40 G49 G80 G90\nT1 M6 (D12 END MILL)\nG54 G0 X0. Y0. S8000 M3\nG43 H1 Z50. M8\nG0 X-8. Y-8.\nZ2.\nG1 Z-5. F600.\nG1 X128. F1500.\nY88.\nX-8.\nY-8.\nG0 Z50.\nT2 M6 (D6.8 DRILL)\nG43 H2 Z50. S3500 M3\nG81 X20. Y20. Z-14. R2. F350.\nX100.\nY60.\nX20.\nG80 G0 Z100. M9\nM30\n%\n')


if __name__ == '__main__':
    doc = flange_dxf(os.path.join(out, 'flange-DN100.dxf'))
    dxf_to_pdf(doc, os.path.join(out, 'flange-DN100.pdf'))
    schedule_xlsx(os.path.join(out, 'delivery-schedule-wk43.xlsx'))
    report_docx(os.path.join(out, 'FAI-report-NW-1042.docx'))
    kickoff_pptx(os.path.join(out, 'kickoff-bracket-revC.pptx'))
    op20_nc(os.path.join(out, 'OP20.nc'))
    print('ok', sorted(os.listdir(out)))
